#!/usr/bin/env python3
"""Prepara a base pública e gera o dashboard Porsche Velocity em HTML único."""

from __future__ import annotations

import csv
import hashlib
import json
import urllib.request
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "data" / "source"
SOURCE = SOURCE_DIR / "base-porsche-sanitizada.xlsx"
PUBLIC_CSV = ROOT / "data" / "porsche_sales_public.csv"
TEMPLATE = ROOT / "src" / "index.template.html"
OUTPUT = ROOT / "index.html"
QUALITY_JSON = ROOT / "data" / "quality-summary.json"

SOURCE_URL = "https://hermes.dio.me/files/assets/8683bed0-cc33-4e06-bca9-04db9c31f9e2.xlsx"
SOURCE_SHA256 = "f874650657ca4e1a99fd4462b297bfc5fd9915e8e693956d9cec47eda2cb2e2d"
AUTHOR = "0Barone (@gabrielnasciimentoleandro-arch)"
BUILD_REFERENCE_DATE = date(2026, 10, 9)

PUBLIC_FIELDS = [
    "model",
    "family",
    "model_year",
    "price",
    "mileage",
    "payment",
    "city",
    "state",
    "status",
    "status_group",
    "powertrain",
]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def ensure_source() -> None:
    SOURCE_DIR.mkdir(parents=True, exist_ok=True)
    if not SOURCE.exists():
        print("Baixando a base oficial do desafio DIO...")
        request = urllib.request.Request(SOURCE_URL, headers={"User-Agent": "PorscheVelocity/1.0"})
        with urllib.request.urlopen(request, timeout=30) as response, SOURCE.open("wb") as target:
            target.write(response.read())
    digest = sha256(SOURCE)
    if digest != SOURCE_SHA256:
        raise ValueError(f"SHA-256 inesperado para a base: {digest}")


def family_for(model: str) -> str:
    family = model.split()[0]
    if family not in {"911", "718", "Cayenne", "Macan", "Taycan", "Panamera"}:
        raise ValueError(f"Família não reconhecida: {model}")
    return family


def powertrain_for(model: str) -> str:
    if model.startswith("Taycan") or model == "Macan Electric":
        return "Electric"
    if "E-Hybrid" in model:
        return "Hybrid"
    return "Combustion"


def status_group_for(status: str) -> str:
    if status == "Delivered":
        return "Delivered"
    if status == "Cancelled":
        return "Cancelled"
    return "In progress"


def load_and_sanitize() -> tuple[list[dict], dict]:
    workbook = load_workbook(SOURCE, data_only=True, read_only=True)
    worksheet = workbook["Sanitized"]
    rows = worksheet.iter_rows(values_only=True)
    headers = list(next(rows))
    required = {
        "sale_id",
        "SaleDateSanitized",
        "customer_name",
        "PorscheModelSanitized",
        "ModelYearSanitized",
        "SalesPriceSanitized",
        "VehicleMileageSanitized",
        "PayMethodSanitized",
        "CitySanitized",
        "StateSanitized",
        "salesperson",
        "DeliveryStatusSanitized",
    }
    missing = required.difference(headers)
    if missing:
        raise ValueError(f"Colunas obrigatórias ausentes: {sorted(missing)}")

    source_records = [dict(zip(headers, values)) for values in rows]
    if len(source_records) != 100:
        raise ValueError(f"Esperados 100 registros, encontrados {len(source_records)}")
    ids = [int(record["sale_id"]) for record in source_records]
    if len(ids) != len(set(ids)):
        raise ValueError("IDs duplicados na fonte")

    clean: list[dict] = []
    valid_dates: list[date] = []
    invalid_dates = 0
    for source_record in source_records:
        model = str(source_record["PorscheModelSanitized"]).strip()
        status = str(source_record["DeliveryStatusSanitized"]).strip()
        raw_date = str(source_record["SaleDateSanitized"]).strip()
        if raw_date == "INVALID":
            invalid_dates += 1
        else:
            valid_dates.append(datetime.strptime(raw_date, "%Y-%m-%d").date())

        record = {
            "model": model,
            "family": family_for(model),
            "model_year": int(source_record["ModelYearSanitized"]),
            "price": round(float(source_record["SalesPriceSanitized"]), 2),
            "mileage": int(float(source_record["VehicleMileageSanitized"])),
            "payment": str(source_record["PayMethodSanitized"]).strip(),
            "city": str(source_record["CitySanitized"]).strip(),
            "state": str(source_record["StateSanitized"]).strip(),
            "status": status,
            "status_group": status_group_for(status),
            "powertrain": powertrain_for(model),
        }
        if record["price"] <= 0 or record["mileage"] < 0:
            raise ValueError("Preço ou quilometragem inválidos após sanitização")
        clean.append(record)

    today = BUILD_REFERENCE_DATE
    metadata = {
        "project": "Porsche Velocity",
        "version": "1.0.0",
        "author": AUTHOR,
        "generated_on": today.isoformat(),
        "source_url": SOURCE_URL,
        "source_sha256": SOURCE_SHA256,
        "records": len(clean),
        "source_unique_ids": len(set(ids)),
        "invalid_dates": invalid_dates,
        "valid_dates": len(valid_dates),
        "valid_date_min": min(valid_dates).isoformat(),
        "valid_date_max": max(valid_dates).isoformat(),
        "future_dates_at_build": sum(value > today for value in valid_dates),
        "public_fields": PUBLIC_FIELDS,
        "removed_fields": [
            "sale_id",
            "customer_name",
            "salesperson",
            "all raw/non-sanitized columns",
            "sale date values",
        ],
        "quality_note": "Datas foram excluídas dos indicadores por conterem valores INVALID e datas futuras.",
    }
    return clean, metadata


def write_public_csv(records: list[dict]) -> None:
    PUBLIC_CSV.parent.mkdir(parents=True, exist_ok=True)
    with PUBLIC_CSV.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=PUBLIC_FIELDS)
        writer.writeheader()
        writer.writerows(records)


def write_quality_summary(records: list[dict], metadata: dict) -> None:
    summary = {
        **metadata,
        "totals": {
            "gross_value": round(sum(record["price"] for record in records), 2),
            "active_value": round(sum(record["price"] for record in records if record["status_group"] != "Cancelled"), 2),
            "active_transactions": sum(record["status_group"] != "Cancelled" for record in records),
            "delivered_transactions": sum(record["status_group"] == "Delivered" for record in records),
            "cancelled_transactions": sum(record["status_group"] == "Cancelled" for record in records),
        },
        "dimensions": {
            "families": dict(sorted(Counter(record["family"] for record in records).items())),
            "states": len(set(record["state"] for record in records)),
            "cities": len(set(record["city"] for record in records)),
            "payment_methods": dict(sorted(Counter(record["payment"] for record in records).items())),
            "status_groups": dict(sorted(Counter(record["status_group"] for record in records).items())),
        },
    }
    QUALITY_JSON.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def generate_html(records: list[dict], metadata: dict) -> None:
    template = TEMPLATE.read_text(encoding="utf-8")
    dataset = json.dumps(records, ensure_ascii=False, separators=(",", ":"))
    build_metadata = json.dumps(metadata, ensure_ascii=False, separators=(",", ":"))
    if "__DATASET_JSON__" not in template or "__BUILD_METADATA_JSON__" not in template:
        raise ValueError("Marcadores de dados ausentes no template")
    html = template.replace("__DATASET_JSON__", dataset).replace("__BUILD_METADATA_JSON__", build_metadata)
    OUTPUT.write_text(html, encoding="utf-8")


def main() -> None:
    ensure_source()
    records, metadata = load_and_sanitize()
    write_public_csv(records)
    write_quality_summary(records, metadata)
    generate_html(records, metadata)

    active = [record for record in records if record["status_group"] != "Cancelled"]
    print(f"Dashboard criado: {OUTPUT}")
    print(f"Base pública: {PUBLIC_CSV}")
    print(f"Registros: {len(records)} | Negócios ativos: {len(active)}")
    print(f"Valor ativo: US$ {sum(record['price'] for record in active):,.2f}")
    print(f"Campos identificáveis removidos: sale_id, customer_name e salesperson")


if __name__ == "__main__":
    main()

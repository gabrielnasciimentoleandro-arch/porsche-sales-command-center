#!/usr/bin/env python3
"""Validação estrutural, financeira, funcional e de privacidade do dashboard."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import struct
import subprocess
import tempfile
from collections import Counter, defaultdict
from html.parser import HTMLParser
from math import isclose
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "source" / "base-porsche-sanitizada.xlsx"
CSV_FILE = ROOT / "data" / "porsche_sales_public.csv"
QUALITY_FILE = ROOT / "data" / "quality-summary.json"
HTML_FILE = ROOT / "index.html"
TEMPLATE = ROOT / "src" / "index.template.html"
README = ROOT / "README.md"
SKILL = ROOT / "skills" / "porsche-dashboard" / "SKILL.md"
CHECKSUMS_FILE = ROOT / "CHECKSUMS.sha256"
SOURCE_SHA256 = "f874650657ca4e1a99fd4462b297bfc5fd9915e8e693956d9cec47eda2cb2e2d"
PUBLIC_FIELDS = {
    "model", "family", "model_year", "price", "mileage", "payment",
    "city", "state", "status", "status_group", "powertrain",
}
REQUIRED_IDS = {
    "dashboard", "familyFilter", "stateFilter", "yearFilter", "paymentFilter",
    "statusFilter", "resetButton", "exportButton", "activeValue", "activeDeals",
    "averageTicket", "deliveryRate", "cancelledValue", "familyChart", "stateChart",
    "paymentChart", "statusDonut", "statusLegend", "portfolioChart", "modelTable",
    "methodologyDialog", "themeButton", "dynamicInsight",
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class DashboardParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.ids: set[str] = set()
        self.external_scripts: list[str] = []
        self.stylesheets: list[str] = []
        self.inline_scripts: list[str] = []
        self._script: list[str] | None = None
        self.lang = ""
        self.title_parts: list[str] = []
        self._in_title = False
        self.metas: list[dict[str, str]] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        values = {key: value or "" for key, value in attrs}
        if tag == "html":
            self.lang = values.get("lang", "")
        if values.get("id"):
            self.ids.add(values["id"])
        if tag == "script":
            if values.get("src"):
                self.external_scripts.append(values["src"])
            else:
                self._script = []
        if tag == "link" and values.get("rel") == "stylesheet":
            self.stylesheets.append(values.get("href", ""))
        if tag == "meta":
            self.metas.append(values)
        if tag == "title":
            self._in_title = True

    def handle_endtag(self, tag: str) -> None:
        if tag == "script" and self._script is not None:
            self.inline_scripts.append("".join(self._script))
            self._script = None
        if tag == "title":
            self._in_title = False

    def handle_data(self, data: str) -> None:
        if self._script is not None:
            self._script.append(data)
        if self._in_title:
            self.title_parts.append(data)


def extract_embedded_data(html: str) -> tuple[list[dict], dict]:
    dataset_match = re.search(r"const SALES = (\[.*?\]);\s*const BUILD = (\{.*?\});", html, re.S)
    if not dataset_match:
        raise AssertionError("Dataset embutido não encontrado")
    return json.loads(dataset_match.group(1)), json.loads(dataset_match.group(2))


def main() -> None:
    required_files = (
        SOURCE, CSV_FILE, QUALITY_FILE, HTML_FILE, TEMPLATE, README, SKILL,
        CHECKSUMS_FILE, ROOT / ".nojekyll",
        ROOT / "docs" / "dashboard-capa.png",
        ROOT / "docs" / "dashboard-filtro-911.png",
        ROOT / "docs" / "dashboard-geral.png",
        ROOT / "docs" / "dashboard-mobile.png",
        ROOT / "docs" / "PROMPT.md",
        ROOT / "docs" / "METODOLOGIA.md",
        ROOT / "docs" / "VALIDACAO.md",
    )
    for path in required_files:
        assert path.exists(), f"Arquivo ausente: {path.relative_to(ROOT)}"
    assert sha256(SOURCE) == SOURCE_SHA256, "A base-fonte foi alterada"

    for checksum_line in CHECKSUMS_FILE.read_text(encoding="utf-8").splitlines():
        expected_hash, relative_path = checksum_line.split(maxsplit=1)
        artifact = ROOT / relative_path.removeprefix("./")
        assert artifact.exists(), f"Artefato do checksum ausente: {relative_path}"
        assert sha256(artifact) == expected_hash, f"Checksum divergente: {relative_path}"

    readme = README.read_text(encoding="utf-8")
    for required_section in (
        "## Perguntas de negócio", "## Tratamento dos dados", "## Prompt e refinamento",
        "## Como a IA foi utilizada", "## Validação", "## Fonte e integridade",
    ):
        assert required_section in readme, f"Seção ausente no README: {required_section}"
    assert "https://gabrielnasciimentoleandro-arch.github.io/porsche-sales-command-center/" in readme
    assert "agente com ferramentas e skill própria" in readme
    assert "ChatGPT Canvas" in readme
    assert "data/source/" in (ROOT / ".gitignore").read_text(encoding="utf-8")

    for image_name, minimum_size in {
        "dashboard-capa.png": (1200, 800),
        "dashboard-filtro-911.png": (1200, 800),
        "dashboard-geral.png": (1200, 2000),
        "dashboard-mobile.png": (390, 800),
    }.items():
        image_bytes = (ROOT / "docs" / image_name).read_bytes()
        assert image_bytes[:8] == b"\x89PNG\r\n\x1a\n"
        width, height = struct.unpack(">II", image_bytes[16:24])
        assert width >= minimum_size[0] and height >= minimum_size[1], image_name

    html = HTML_FILE.read_text(encoding="utf-8")
    assert 35_000 < len(html.encode("utf-8")) < 1_000_000
    assert "__DATASET_JSON__" not in html
    assert "__BUILD_METADATA_JSON__" not in html
    assert "https://cdn" not in html.lower()
    assert "fetch(" not in html
    assert "XMLHttpRequest" not in html
    assert "WebSocket" not in html
    assert "connect-src 'none'" in html
    assert "@media (max-width: 760px)" in html
    assert html.count("Pergunta 0") >= 4

    parser = DashboardParser()
    parser.feed(html)
    assert parser.lang == "pt-BR"
    assert "Porsche Velocity" in "".join(parser.title_parts)
    assert not parser.external_scripts, f"Scripts externos: {parser.external_scripts}"
    assert not parser.stylesheets, f"CSS externo: {parser.stylesheets}"
    assert REQUIRED_IDS.issubset(parser.ids), f"IDs ausentes: {sorted(REQUIRED_IDS - parser.ids)}"
    assert any(meta.get("name") == "viewport" for meta in parser.metas)
    assert any(meta.get("http-equiv") == "Content-Security-Policy" for meta in parser.metas)
    assert len(parser.inline_scripts) == 1

    records, metadata = extract_embedded_data(html)
    assert len(records) == 100
    assert metadata["records"] == 100
    assert metadata["source_sha256"] == SOURCE_SHA256
    assert metadata["invalid_dates"] == 24
    assert metadata["valid_dates"] == 76
    assert metadata["valid_date_min"] == "2024-03-14"
    assert metadata["valid_date_max"] == "2027-10-01"
    assert set(metadata["public_fields"]) == PUBLIC_FIELDS
    assert all(set(record) == PUBLIC_FIELDS for record in records)

    with CSV_FILE.open("r", encoding="utf-8", newline="") as handle:
        csv_rows = list(csv.DictReader(handle))
    assert len(csv_rows) == 100
    assert set(csv_rows[0]) == PUBLIC_FIELDS
    for csv_row, record in zip(csv_rows, records):
        normalized = {
            **csv_row,
            "model_year": int(csv_row["model_year"]),
            "price": float(csv_row["price"]),
            "mileage": int(csv_row["mileage"]),
        }
        assert normalized == record

    # Confere que os campos públicos foram derivados exatamente das colunas sanitizadas.
    workbook = load_workbook(SOURCE, data_only=True, read_only=True)
    worksheet = workbook["Sanitized"]
    iterator = worksheet.iter_rows(values_only=True)
    headers = list(next(iterator))
    source_rows = [dict(zip(headers, values)) for values in iterator]
    assert len(source_rows) == 100
    assert len({row["sale_id"] for row in source_rows}) == 100
    for source, public in zip(source_rows, records):
        assert public["model"] == source["PorscheModelSanitized"]
        assert public["model_year"] == int(source["ModelYearSanitized"])
        assert isclose(public["price"], float(source["SalesPriceSanitized"]), abs_tol=1e-9)
        assert public["mileage"] == int(float(source["VehicleMileageSanitized"]))
        assert public["payment"] == source["PayMethodSanitized"]
        assert public["city"] == source["CitySanitized"]
        assert public["state"] == source["StateSanitized"]
        assert public["status"] == source["DeliveryStatusSanitized"]

    # Privacidade: nenhum valor de nome de cliente ou vendedor entra no dataset embutido.
    public_json = json.dumps(records, ensure_ascii=False).casefold()
    for field in ("customer_name", "salesperson"):
        for value in {str(row[field]).strip() for row in source_rows}:
            if len(value) >= 4:
                assert value.casefold() not in public_json, f"Valor identificável encontrado: {value}"
    forbidden_fields = {
        "sale_id", "sale_date", "customer_name", "salesperson", "porsche_model",
        "sale_price", "vehicle_mileage", "payment_method", "delivery_status",
    }
    assert forbidden_fields.isdisjoint(records[0])

    gross = sum(record["price"] for record in records)
    active = [record for record in records if record["status_group"] != "Cancelled"]
    delivered = [record for record in records if record["status_group"] == "Delivered"]
    cancelled = [record for record in records if record["status_group"] == "Cancelled"]
    assert isclose(gross, 12_827_800.5, abs_tol=1e-6)
    assert len(active) == 93
    assert isclose(sum(record["price"] for record in active), 12_202_550.5, abs_tol=1e-6)
    assert len(delivered) == 41
    assert isclose(sum(record["price"] for record in delivered), 5_078_800.5, abs_tol=1e-6)
    assert len(cancelled) == 7
    assert isclose(sum(record["price"] for record in cancelled), 625_250.0, abs_tol=1e-6)

    assert Counter(record["family"] for record in records) == {
        "911": 23, "718": 12, "Cayenne": 18, "Macan": 17, "Taycan": 16, "Panamera": 14,
    }
    assert Counter(record["status_group"] for record in records) == {
        "Delivered": 41, "In progress": 52, "Cancelled": 7,
    }
    assert Counter(record["powertrain"] for record in active) == {
        "Combustion": 68, "Electric": 20, "Hybrid": 5,
    }

    family_value: defaultdict[str, float] = defaultdict(float)
    state_value: defaultdict[str, float] = defaultdict(float)
    payment_value: defaultdict[str, float] = defaultdict(float)
    for record in active:
        family_value[record["family"]] += record["price"]
        state_value[record["state"]] += record["price"]
        payment_value[record["payment"]] += record["price"]
    assert max(family_value, key=family_value.get) == "911"
    assert isclose(family_value["911"], 4_168_100.0, abs_tol=1e-6)
    assert max(state_value, key=state_value.get) == "TX"
    assert isclose(state_value["TX"], 2_023_400.5, abs_tol=1e-6)
    assert max(payment_value, key=payment_value.get) == "Wire Transfer"
    assert isclose(payment_value["Wire Transfer"], 3_979_450.0, abs_tol=1e-6)

    quality = json.loads(QUALITY_FILE.read_text(encoding="utf-8"))
    assert quality["totals"]["gross_value"] == 12_827_800.5
    assert quality["totals"]["active_value"] == 12_202_550.5
    assert quality["totals"]["active_transactions"] == 93

    # Node verifica a sintaxe do JavaScript embutido quando disponível.
    node = shutil.which("node")
    if node:
        with tempfile.NamedTemporaryFile("w", suffix=".js", encoding="utf-8", delete=False) as script_file:
            script_file.write(parser.inline_scripts[0])
            script_path = Path(script_file.name)
        try:
            result = subprocess.run([node, "--check", str(script_path)], capture_output=True, text=True)
            assert result.returncode == 0, result.stderr
        finally:
            script_path.unlink(missing_ok=True)

    print("VALIDAÇÃO APROVADA")
    print("HTML único: CSS, JavaScript e 100 registros incorporados")
    print("Interatividade: 5 filtros, reset, exportação CSV, tema e filtros por clique")
    print("Perguntas de negócio: portfólio, mercados, pagamento e fluxo de entrega")
    print("Valor bruto: US$ 12.827.800,50")
    print("Valor ativo: US$ 12.202.550,50 | 93 negócios")
    print("Entregues: 41 | Em andamento: 52 | Cancelados: 7")
    print("Privacidade: IDs, clientes, vendedores, colunas cruas e datas ausentes do HTML")
    print("Dependências externas e chamadas de rede: AUSENTES")
    print(f"SHA-256 da fonte: {SOURCE_SHA256}")


if __name__ == "__main__":
    main()
